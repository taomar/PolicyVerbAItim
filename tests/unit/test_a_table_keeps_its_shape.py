"""Guard: a grid's shape survives from the parser to whatever finally reads it.

WHAT THIS PROTECTS
------------------
A table states part of its meaning in position. "5" under "Days" and "5" under
"Occurrences" are different facts, and the only thing that tells them apart is
which column the value sits in. Both parsers in this platform can see that, and
both used to throw it away by the time anything downstream looked:

  * the row parser joins a row's cells with a separator that is nowhere in the
    source and passes on the joined line;
  * the cell parser keeps coordinates in memory and drops them at the clause
    projection, so a document rebuilt from storage -- which is the only kind
    anything downstream ever sees -- had no cell to build an edge from.

Either way the shape reached storage and stopped. The only route back was to
split the rendered line on the separator, which is a guess: it assumes no cell
contains the separator, and nothing establishes that.

These tests pin the carrier that removes the need to guess. The cells a parser
read, each with the row, column and span it was read at, travel with the element
through the clause projection, through storage, back through the rebuild, and
into what a reader is shown.

WHAT IS DELIBERATELY NOT CLAIMED
--------------------------------
Nothing here derives a shape that was not recorded. A row stored before the
carrier existed has no cells and is given none; a column nobody named stays
unnamed rather than being handed the nearest free label. Absence travels as
absence, which is the difference between a record and an assumption -- and it is
asserted as explicitly as the presence is.

Written against invented grids of several shapes, with a vocabulary that names
nothing: no document, no domain, no language, no observed count. The last test
in this module holds the code that implements the carrier to the same rule.
"""
from __future__ import annotations

import ast
import inspect
import re
import textwrap

import pytest

from policy_platform.contracts.canonical_document import (
    TABLE_STRUCTURE_KEY,
    CanonicalDocument,
    CanonicalElement,
    CanonicalPage,
    SourceFragment,
    TableCell,
    TableCellRef,
    TableStructure,
    covered_columns,
    table_cell_of,
    table_structure_of,
)
from policy_platform.contracts.reading_plan import (
    build_reading_plan,
    render_table_cells,
    render_table_columns,
)
from policy_platform.contracts.structural_graph import build_structural_graph
from policy_platform.domain.models import Clause
from policy_platform.infrastructure.extraction import ai_extraction
from policy_platform.infrastructure.ingestion import canonical_rebuild, document_ingestion
from policy_platform.infrastructure.ingestion.canonical_fidelity import (
    rebuild_row_text,
    verify_element_text,
)
from policy_platform.infrastructure.ingestion.canonical_rebuild import (
    canonical_from_clauses,
    stored_fragments,
    stored_table_structure,
)
from policy_platform.infrastructure.ingestion.document_extraction import (
    ClauseData,
    clauses_from_document,
)

#: Grid shapes, so nothing comes to depend on a table having a particular width,
#: a particular number of rows, or more rows than columns.
GRID_SHAPES = [(1, 2), (1, 5), (3, 3), (2, 4), (4, 2)]

#: The edge kinds that exist only because a cell knows where it sits.
TABLE_EDGE_KINDS = ("table_cell_of", "header_for", "merged_with")

#: The separator a row's cells are joined with when a row becomes a line.
ROW_SEPARATOR = " | "


# ── fixtures: invented grids, no domain anywhere ─────────────────────────


def _label(column: int) -> str:
    return f"Col{column}"


def _value(row: int, column: int) -> str:
    return f"v{row}-{column}"


def _fragment(text: str, offset: int, page: int = 1) -> SourceFragment:
    return SourceFragment(
        page=page, start_offset=offset, end_offset=offset + len(text), text=text
    )


def _row_element(
    *,
    element_id: str,
    order: int,
    row_index: int,
    cells: list[str],
    headers: list[str] | None,
    table_id: str = "t1",
    offset: int = 0,
    page: int = 1,
) -> CanonicalElement:
    """One element holding a whole row, the shape a row parser produces."""

    text = ROW_SEPARATOR.join(cells)
    return CanonicalElement(
        element_id=element_id,
        element_type="table_row",
        logical_order=order,
        text=text,
        table_id=table_id,
        table_headers=headers,
        transformations=["table_cell_join"],
        source_fragments=[_fragment(text, offset, page)],
        table_structure=TableStructure(
            cells=[
                TableCell(
                    text=value,
                    position=TableCellRef(row_index=row_index, column_index=column),
                )
                for column, value in enumerate(cells)
            ],
            column_labels=headers,
        ),
    )


def _row_document(
    body_rows: int, columns: int, *, headers: list[str] | None
) -> CanonicalDocument:
    elements: list[CanonicalElement] = []
    offset = 0
    for row in range(body_rows):
        cells = [_value(row, column) for column in range(columns)]
        element = _row_element(
            element_id=f"E{row:06d}",
            order=row,
            # Row 0 of the grid is the header the labels came from, so the body
            # starts at 1 -- the index a reader counting rows on the page sees.
            row_index=row + 1,
            cells=cells,
            headers=headers,
            offset=offset,
        )
        offset += len(element.text)
        elements.append(element)

    raw = "".join(element.source_fragments[0].text for element in elements)
    return CanonicalDocument(
        document_id="rows",
        page_count=1,
        pages=[CanonicalPage(page=1, raw_text=raw)],
        elements=elements,
        parser="test",
    )


def _cell_document(body_rows: int, columns: int) -> CanonicalDocument:
    """A grid articulated one element per cell, the shape a cell parser produces.

    Only ``table_cell`` is set on each element -- no structure. That is what a
    cell parser produces today, and it is the input the carrier has to lift, so
    building it any other way would test the fixture rather than the code.
    """

    elements: list[CanonicalElement] = []
    order = 0
    offset = 0
    for row in range(body_rows + 1):
        for column in range(columns):
            is_header = row == 0
            text = _label(column) if is_header else _value(row, column)
            elements.append(
                CanonicalElement(
                    element_id=f"E{order:06d}",
                    element_type="table_cell",
                    logical_order=order,
                    text=text,
                    table_id="t1",
                    table_cell=TableCellRef(
                        row_index=row, column_index=column, is_header=is_header
                    ),
                    source_fragments=[_fragment(text, offset)],
                )
            )
            offset += len(text)
            order += 1

    raw = "".join(
        fragment.text for element in elements for fragment in element.source_fragments
    )
    return CanonicalDocument(
        document_id="cells",
        page_count=1,
        pages=[CanonicalPage(page=1, raw_text=raw)],
        elements=elements,
        parser="test",
    )


def _stored(document: CanonicalDocument) -> list[Clause]:
    """Real ORM rows, built the way the repository builds them.

    The ORM model rather than a stand-in, because the claim under test is that
    the shape reaches *storage*: a stand-in with hand-written attributes would
    pass whether or not the row has anywhere to put it.
    """

    return [
        Clause(
            clause_ref=data.clause_ref,
            section=data.section,
            page=data.page,
            text=data.text,
            sequence=index,
            element_id=data.element_id,
            element_type=data.element_type,
            source_fragments=data.source_fragments,
            table_id=data.table_id,
            table_headers=data.table_headers,
        )
        for index, data in enumerate(clauses_from_document(document))
    ]


def _round_trip(document: CanonicalDocument) -> CanonicalDocument:
    return canonical_from_clauses(document.document_id, _stored(document))


def _strip_structure(clauses: list[Clause]) -> list[Clause]:
    """The same rows as they were stored before the carrier existed."""

    for clause in clauses:
        clause.source_fragments = [
            entry
            for entry in (clause.source_fragments or [])
            if TABLE_STRUCTURE_KEY not in entry
        ]
    return clauses


# ── the carrier itself ───────────────────────────────────────────────────


class TestTheCarrierReadsItsOwnCoordinates:
    """The record answers questions about a grid without consulting any text."""

    def test_cells_come_back_in_reading_order_however_they_went_in(self) -> None:
        """Order is a property of the coordinates, not of the insertion loop."""

        positions = [(1, 1), (0, 2), (1, 0), (0, 0), (1, 2), (0, 1)]
        structure = TableStructure(
            cells=[
                TableCell(
                    text=_value(row, column),
                    position=TableCellRef(row_index=row, column_index=column),
                )
                for row, column in positions
            ]
        )

        assert [
            (cell.position.row_index, cell.position.column_index)
            for cell in structure.ordered_cells
        ] == sorted(positions)
        assert [row for row, _ in structure.rows] == [0, 1]
        assert [len(cells) for _, cells in structure.rows] == [3, 3]

    def test_a_repeated_value_stays_as_many_cells_as_it_was_written(self) -> None:
        """Identical text in several columns is several facts, not one.

        A grid may legitimately state the same value under three headings. A
        record that folded them together would report one, and the reader could
        never discover which columns were covered.
        """

        repeated = "same"
        structure = TableStructure(
            cells=[
                TableCell(
                    text=repeated,
                    position=TableCellRef(row_index=1, column_index=column),
                )
                for column in range(3)
            ],
            column_labels=[_label(column) for column in range(3)],
        )

        pairs = structure.header_value_pairs()
        assert [cell.text for _, cell in pairs] == [repeated] * 3
        assert [names for names, _ in pairs] == [(_label(c),) for c in range(3)]

    def test_a_merged_header_names_every_column_it_covers(self) -> None:
        """A banner over three columns qualifies the value in each of them.

        Filing it under the column it starts in would leave the other two
        unnamed while the source plainly names them, and would attribute a claim
        that holds across a band to one column of it.
        """

        banner = "Band"
        structure = TableStructure(
            cells=[
                TableCell(
                    text=banner,
                    position=TableCellRef(
                        row_index=0, column_index=0, column_span=3, is_header=True
                    ),
                ),
                *[
                    TableCell(
                        text=_value(1, column),
                        position=TableCellRef(row_index=1, column_index=column),
                    )
                    for column in range(3)
                ],
            ]
        )

        assert structure.labels_by_column() == {column: (banner,) for column in range(3)}
        assert [names for names, _ in structure.header_value_pairs()] == [(banner,)] * 3

    def test_a_value_spanning_columns_carries_every_name_above_it(self) -> None:
        """The span cuts the other way too, and both names are kept in order."""

        structure = TableStructure(
            cells=[
                TableCell(
                    text=_label(0),
                    position=TableCellRef(row_index=0, column_index=0, is_header=True),
                ),
                TableCell(
                    text=_label(1),
                    position=TableCellRef(row_index=0, column_index=1, is_header=True),
                ),
                TableCell(
                    text="wide",
                    position=TableCellRef(row_index=1, column_index=0, column_span=2),
                ),
            ]
        )

        names, cell = structure.header_value_pairs()[0]
        assert cell.text == "wide"
        assert names == (_label(0), _label(1))

    def test_a_stacked_header_keeps_both_of_its_rows_outermost_first(self) -> None:
        """Two header rows are two names for one column, and order is meaning."""

        structure = TableStructure(
            cells=[
                TableCell(
                    text="Outer",
                    position=TableCellRef(
                        row_index=0, column_index=0, column_span=2, is_header=True
                    ),
                ),
                TableCell(
                    text="Inner",
                    position=TableCellRef(row_index=1, column_index=1, is_header=True),
                ),
                TableCell(
                    text=_value(2, 1),
                    position=TableCellRef(row_index=2, column_index=1),
                ),
            ]
        )

        assert structure.labels_by_column()[1] == ("Outer", "Inner")

    def test_a_cell_covered_by_a_merge_is_not_reported_as_a_second_value(self) -> None:
        """A placeholder points at the value it belongs to; it is not one."""

        owner = "E-owner"
        structure = TableStructure(
            cells=[
                TableCell(
                    text="held",
                    element_id=owner,
                    position=TableCellRef(row_index=1, column_index=0, column_span=2),
                ),
                TableCell(
                    text="held",
                    position=TableCellRef(
                        row_index=1, column_index=1, merged_into=owner
                    ),
                ),
            ],
            column_labels=[_label(0), _label(1)],
        )

        pairs = structure.header_value_pairs()
        assert [cell.element_id for _, cell in pairs] == [owner]
        # The placeholder is still in the record: it is part of the grid's shape
        # and only the *value* reading skips it.
        assert len(structure.ordered_cells) == 2

    def test_a_column_nobody_named_stays_unnamed(self) -> None:
        """Absent is not empty, and it is certainly not the next label along."""

        structure = TableStructure(
            cells=[
                TableCell(
                    text=_value(1, column),
                    position=TableCellRef(row_index=1, column_index=column),
                )
                for column in range(3)
            ],
            # Shorter than the row, and blank in the middle: two ways a grid
            # legitimately fails to name a column.
            column_labels=[_label(0), "  "],
        )

        assert structure.labels_by_column() == {0: (_label(0),)}
        assert [names for names, _ in structure.header_value_pairs()] == [
            (_label(0),),
            (),
            (),
        ]

    def test_a_grid_that_named_nothing_pairs_nothing(self) -> None:
        """No labels recorded means no labels reported, for every shape."""

        structure = TableStructure(
            cells=[
                TableCell(
                    text=_value(1, column),
                    position=TableCellRef(row_index=1, column_index=column),
                )
                for column in range(4)
            ]
        )

        assert structure.labels_by_column() == {}
        assert all(names == () for names, _ in structure.header_value_pairs())

    def test_a_header_cell_outranks_a_positional_label_for_its_own_columns(
        self,
    ) -> None:
        """A recorded coordinate beats an assumed index, and only where it reaches."""

        structure = TableStructure(
            cells=[
                TableCell(
                    text="Stated",
                    position=TableCellRef(row_index=0, column_index=0, is_header=True),
                ),
                TableCell(text=_value(1, 0), position=TableCellRef(row_index=1, column_index=0)),
                TableCell(text=_value(1, 1), position=TableCellRef(row_index=1, column_index=1)),
            ],
            column_labels=["Assumed0", "Assumed1"],
        )

        assert structure.labels_by_column() == {0: ("Stated",), 1: ("Assumed1",)}

    def test_a_span_covers_every_column_it_reaches(self) -> None:
        """One rule for what a span covers, shared by everything that resolves one."""

        assert list(covered_columns(TableCellRef(row_index=0, column_index=2))) == [2]
        assert list(
            covered_columns(TableCellRef(row_index=0, column_index=2, column_span=3))
        ) == [2, 3, 4]


class TestWhichElementsHaveAShapeAtAll:
    """The bridge between the two parse shapes, and its refusals."""

    def test_a_cell_element_lends_its_own_coordinate_to_the_carrier(self) -> None:
        element = _cell_document(1, 2).elements[0]
        structure = table_structure_of(element)

        assert structure is not None
        assert structure.sole_cell is not None
        assert structure.sole_cell.position == element.table_cell
        assert structure.sole_cell.text == element.text
        assert structure.sole_cell.element_id == element.element_id

    def test_prose_is_given_no_shape(self) -> None:
        """An element that was never part of a grid does not acquire one."""

        text = "A paragraph that is part of no grid."
        element = CanonicalElement(
            element_id="E000000",
            element_type="paragraph",
            logical_order=0,
            text=text,
            source_fragments=[_fragment(text, 0)],
        )

        assert table_structure_of(element) is None
        assert table_cell_of(element.element_type, None) is None

    def test_only_a_single_cell_element_takes_a_coordinate_from_its_shape(self) -> None:
        """A row does not sit in a column, so a row is given no column."""

        row = _row_element(
            element_id="E000000",
            order=0,
            row_index=1,
            cells=[_value(1, 0), _value(1, 1)],
            headers=None,
        )
        structure = table_structure_of(row)

        assert structure is not None
        assert structure.sole_cell is None
        assert table_cell_of("table_row", structure) is None
        assert table_cell_of("table_cell", structure) is None


# ── continuity: extraction -> storage -> rebuild ─────────────────────────


class TestTheShapeSurvivesStorage:
    @pytest.mark.parametrize(("body_rows", "columns"), GRID_SHAPES)
    def test_a_rows_cells_come_back_exactly_as_they_went_in(
        self, body_rows: int, columns: int
    ) -> None:
        """Values, coordinates, spans and order, all unchanged by the trip."""

        headers = [_label(column) for column in range(columns)]
        document = _row_document(body_rows, columns, headers=headers)
        rebuilt = _round_trip(document)

        for before, after in zip(document.elements, rebuilt.elements):
            assert after.table_structure is not None
            assert before.table_structure is not None
            assert after.table_structure.model_dump() == before.table_structure.model_dump()
            assert [cell.text for cell in after.table_structure.ordered_cells] == [
                cell.text for cell in before.table_structure.ordered_cells
            ]

    @pytest.mark.parametrize(("body_rows", "columns"), GRID_SHAPES)
    def test_the_row_index_a_reader_would_count_survives(
        self, body_rows: int, columns: int
    ) -> None:
        """Row identity is ordered and stable, not re-derived from list position."""

        document = _row_document(body_rows, columns, headers=None)
        rebuilt = _round_trip(document)

        indexes = [
            element.table_structure.ordered_cells[0].position.row_index
            for element in rebuilt.elements
            if element.table_structure is not None
        ]
        assert indexes == list(range(1, body_rows + 1))

    @pytest.mark.parametrize(("body_rows", "columns"), GRID_SHAPES)
    def test_a_cell_parse_keeps_its_coordinates_through_storage(
        self, body_rows: int, columns: int
    ) -> None:
        """The other parse shape, carried by the same record."""

        document = _cell_document(body_rows, columns)
        rebuilt = _round_trip(document)

        assert [element.table_cell for element in rebuilt.elements] == [
            element.table_cell for element in document.elements
        ]
        assert all(
            element.table_structure is not None for element in rebuilt.elements
        )

    @pytest.mark.parametrize(("body_rows", "columns"), GRID_SHAPES)
    def test_the_pairing_a_rebuilt_row_reports_is_the_one_it_was_parsed_with(
        self, body_rows: int, columns: int
    ) -> None:
        """The question the whole carry exists to answer, asked after storage."""

        headers = [_label(column) for column in range(columns)]
        rebuilt = _round_trip(_row_document(body_rows, columns, headers=headers))

        for element in rebuilt.elements:
            assert element.table_structure is not None
            pairs = element.table_structure.header_value_pairs()
            assert [names for names, _ in pairs] == [(h,) for h in headers]
            assert [cell.position.column_index for _, cell in pairs] == list(
                range(columns)
            )

    def test_a_merge_survives_storage_with_its_span_intact(self) -> None:
        """The span is the qualification; losing it narrows what the grid said."""

        banner = "Band"
        element = CanonicalElement(
            element_id="E000000",
            element_type="table_row",
            logical_order=0,
            text="wide",
            table_id="t1",
            transformations=["table_cell_join"],
            source_fragments=[_fragment("wide", 0)],
            table_structure=TableStructure(
                cells=[
                    TableCell(
                        text=banner,
                        position=TableCellRef(
                            row_index=0, column_index=0, column_span=2, is_header=True
                        ),
                    ),
                    TableCell(
                        text="wide",
                        position=TableCellRef(
                            row_index=1, column_index=0, column_span=2, row_span=2
                        ),
                    ),
                ]
            ),
        )
        document = CanonicalDocument(
            document_id="merge",
            page_count=1,
            pages=[CanonicalPage(page=1, raw_text="wide")],
            elements=[element],
            parser="test",
        )

        structure = _round_trip(document).elements[0].table_structure
        assert structure is not None
        positions = [cell.position for cell in structure.ordered_cells]
        assert [p.column_span for p in positions] == [2, 2]
        assert [p.row_span for p in positions] == [1, 2]
        assert structure.header_value_pairs() == [
            ((banner,), structure.ordered_cells[1])
        ]


# ── the graph and the plan, on a document nothing re-parsed ──────────────


class TestWhatTheShapeReaches:
    @pytest.mark.parametrize(("body_rows", "columns"), GRID_SHAPES)
    def test_a_rebuilt_cell_grid_builds_the_same_edges_it_was_parsed_with(
        self, body_rows: int, columns: int
    ) -> None:
        document = _cell_document(body_rows, columns)

        def table_edges(graph) -> list[tuple[str, str, str]]:
            return sorted(
                (edge.source, edge.target, edge.kind)
                for edge in graph.edges
                if edge.kind in TABLE_EDGE_KINDS
            )

        before = table_edges(build_structural_graph(document))
        after = table_edges(build_structural_graph(_round_trip(document)))

        assert before
        assert after == before

    @pytest.mark.parametrize(("body_rows", "columns"), GRID_SHAPES)
    def test_a_reader_is_shown_the_value_under_each_column(
        self, body_rows: int, columns: int
    ) -> None:
        """Every value and every name it was recorded under, and nothing else."""

        headers = [_label(column) for column in range(columns)]
        rebuilt = _round_trip(_row_document(body_rows, columns, headers=headers))
        element = rebuilt.elements[0]
        assert element.table_structure is not None

        line = render_table_cells(element.table_structure)

        for column in range(columns):
            assert f"{_label(column)}: {_value(0, column)}" in line
        assert ROW_SEPARATOR not in line

    def test_a_value_whose_column_was_never_named_is_shown_alone(self) -> None:
        """No borrowed label, and no sentence about this system's own gaps.

        The row is still rendered because *some* column was named; only the
        unnamed value goes bare.
        """

        structure = TableStructure(
            cells=[
                TableCell(text="alpha", position=TableCellRef(row_index=1, column_index=0)),
                TableCell(text="beta", position=TableCellRef(row_index=1, column_index=1)),
            ],
            column_labels=[_label(0)],
        )

        line = render_table_cells(structure)

        assert f"{_label(0)}: alpha" in line
        assert "beta" in line
        assert f"{_label(0)}: beta" not in line

    def test_nothing_is_rendered_for_an_element_with_no_shape(self) -> None:
        assert render_table_cells(None) == ""
        assert render_table_cells(TableStructure()) == ""
        assert render_table_cells(
            TableStructure(
                cells=[
                    TableCell(text="  ", position=TableCellRef(row_index=0, column_index=0))
                ]
            )
        ) == ""

    def test_nothing_is_rendered_when_no_column_was_ever_named(self) -> None:
        """A line that can state no pairing has nothing the value does not.

        Printing the values back with no names attached would take the place of
        the column-names line while saying strictly less than it, which is how a
        richer record makes a reader worse informed.
        """

        assert (
            render_table_cells(
                TableStructure(
                    cells=[
                        TableCell(
                            text=_value(1, column),
                            position=TableCellRef(row_index=1, column_index=column),
                        )
                        for column in range(3)
                    ]
                )
            )
            == ""
        )

    def test_a_single_cell_element_keeps_the_line_it_had(self) -> None:
        """The cell-per-element parse must not lose its column names.

        A cell's shape says where it sits, not what its neighbours are called, so
        pairing has nothing to state and the marker falls through to the names
        the grid recorded. Asserted because the richer record would otherwise
        have replaced a useful line with the cell's own text repeated back.
        """

        document = _cell_document(1, 3)
        clause = next(
            clause
            for clause in _stored(document)
            if clause.element_type == "table_cell"
            and stored_table_structure(clause.source_fragments) is not None
        )
        clause.table_headers = [_label(column) for column in range(3)]

        assert ai_extraction._column_marker(clause) == render_table_columns(
            clause.table_headers
        )
        assert ai_extraction._column_marker(clause) != ""

    def test_the_batch_marker_prefers_the_pairing_and_falls_back_without_it(
        self,
    ) -> None:
        """The one place the two lines meet, and the rule for choosing.

        A clause whose shape was recorded is described cell by cell. A clause
        stored before the carrier -- the same row, with the record taken away --
        gets exactly the line it got before, with no coordinate invented for it.
        """

        headers = [_label(0), _label(1)]
        document = _row_document(1, 2, headers=headers)
        structured, legacy = _stored(document)[0], _strip_structure(_stored(document))[0]

        assert ai_extraction._column_marker(structured) == render_table_cells(
            document.elements[0].table_structure
        )
        assert ai_extraction._column_marker(legacy) == render_table_columns(headers)
        assert ai_extraction._column_marker(structured) != ""

    def test_a_reading_plan_frames_a_rebuilt_cell_with_its_header(self) -> None:
        """The claim that decides whether a shape reaches a model at all."""

        document = _cell_document(2, 3)
        rebuilt = _round_trip(document)
        graph = build_structural_graph(rebuilt)
        plan = build_reading_plan(rebuilt, graph)

        framed = {
            target
            for unit in plan.units
            for target in unit.target_element_ids
            if graph.sources(target, "header_for")
        }
        assert framed


# ── compatibility: records written before any of this existed ────────────


class TestALegacyRecordIsUnchanged:
    def test_a_clause_with_no_shape_rebuilds_exactly_as_before(self) -> None:
        """Absent stays absent, and nothing else about the row moves."""

        document = _row_document(2, 3, headers=[_label(0), _label(1), _label(2)])
        legacy = _strip_structure(_stored(document))

        rebuilt = canonical_from_clauses(document.document_id, legacy)

        assert all(element.table_structure is None for element in rebuilt.elements)
        assert all(element.table_cell is None for element in rebuilt.elements)
        assert [element.text for element in rebuilt.elements] == [
            element.text for element in document.elements
        ]
        assert [element.table_headers for element in rebuilt.elements] == [
            element.table_headers for element in document.elements
        ]
        assert [
            [fragment.model_dump() for fragment in element.source_fragments]
            for element in rebuilt.elements
        ] == [
            [fragment.model_dump() for fragment in element.source_fragments]
            for element in document.elements
        ]

    def test_prose_is_stored_byte_for_byte_as_it_was(self) -> None:
        """The plain-text consumer sees no difference at all.

        The carrier rides in the clause's provenance list, so the test that
        matters for everything that is not a table is that the list is
        untouched: same entries, same order, same keys.
        """

        text = "An ordinary sentence, belonging to no grid."
        document = CanonicalDocument(
            document_id="prose",
            page_count=1,
            pages=[CanonicalPage(page=1, raw_text=text)],
            elements=[
                CanonicalElement(
                    element_id="E000000",
                    element_type="paragraph",
                    logical_order=0,
                    text=text,
                    source_fragments=[_fragment(text, 0)],
                )
            ],
            parser="test",
        )

        stored = clauses_from_document(document)[0]

        assert stored.source_fragments == [
            document.elements[0].source_fragments[0].model_dump()
        ]
        assert stored_table_structure(stored.source_fragments) is None
        assert stored_fragments(stored.source_fragments) == stored.source_fragments

    def test_an_element_serialised_before_the_field_existed_still_loads(self) -> None:
        """A stored canonical element with no such key reads back as no shape."""

        legacy = {
            "element_id": "E000000",
            "element_type": "table_row",
            "logical_order": 0,
            "text": "a | b",
            "table_id": "t1",
            "table_headers": [_label(0), _label(1)],
            "source_fragments": [
                {"page": 1, "start_offset": 0, "end_offset": 5, "text": "a | b"}
            ],
        }

        element = CanonicalElement.model_validate(legacy)

        assert element.table_structure is None
        assert table_structure_of(element) is None
        # And the field is omitted again on the way out, so a consumer reading
        # the dump sees the shape it has always seen.
        assert "table_structure" not in element.model_dump(exclude_none=True)

    def test_the_shape_round_trips_through_plain_json_shapes(self) -> None:
        """Storage is JSON, so the record has to survive being one."""

        original = _row_document(1, 3, headers=[_label(c) for c in range(3)]).elements[0]
        assert original.table_structure is not None

        payload = original.table_structure.model_dump(exclude_none=True)
        restored = TableStructure.model_validate(payload)

        assert restored.model_dump() == original.table_structure.model_dump()

    def test_a_record_from_another_capability_is_stepped_over(self) -> None:
        """The provenance list is shared, and a reader here owns only its own key.

        Recognised by shape rather than by name, so a namespaced record this
        module has never heard of is skipped instead of being handed to the
        fragment constructor -- which is what would happen if the rule were "skip
        the one key I know about".
        """

        document = _row_document(1, 2, headers=None)
        clause = _stored(document)[0]
        clause.source_fragments = [
            *clause.source_fragments,
            {"__some_other_record__": {"anything": True}},
        ]

        rebuilt = canonical_from_clauses(document.document_id, [clause])
        element = rebuilt.elements[0]

        assert [fragment.text for fragment in element.source_fragments] == [
            document.elements[0].text
        ]
        assert element.table_structure is not None
        assert len(stored_fragments(clause.source_fragments)) == 1

    def test_a_malformed_shape_is_refused_rather_than_read_as_absent(self) -> None:
        """A corrupt record must not become indistinguishable from an honest gap."""

        with pytest.raises(Exception):
            stored_table_structure([{TABLE_STRUCTURE_KEY: {"cells": [{"text": "x"}]}}])

    def test_every_field_the_flatten_produces_still_has_a_column_to_land_in(
        self,
    ) -> None:
        """The carrier travels in an existing column, and adds no new field.

        Stated here as well as in the module that first made the claim, because
        this is the change that would most easily have broken it: a new value to
        persist is exactly when somebody adds a field with nowhere to go.
        """

        import dataclasses

        produced = {field.name for field in dataclasses.fields(ClauseData)}
        assert produced <= set(Clause.__table__.columns.keys())


# ── the fidelity chain ───────────────────────────────────────────────────


class TestARowCanNowBeProved:
    """The link the fidelity module said reconstruction could not close."""

    def _document(self, element: CanonicalElement) -> CanonicalDocument:
        return CanonicalDocument(
            document_id="fidelity",
            page_count=1,
            pages=[CanonicalPage(page=1, raw_text=element.text)],
            elements=[element],
            parser="test",
        )

    def test_a_row_that_records_its_cells_is_verified(self) -> None:
        element = _row_document(1, 3, headers=None).elements[0]

        report = verify_element_text(self._document(element))

        assert report.failures == []
        assert report.unprovable == []
        assert report.verified == 1

    def test_a_row_that_records_none_is_still_unprovable(self) -> None:
        """The older shape keeps the older verdict, which was honest about itself."""

        text = ROW_SEPARATOR.join([_value(1, 0), _value(1, 1)])
        element = CanonicalElement(
            element_id="E000000",
            element_type="table_row",
            logical_order=0,
            text=text,
            transformations=["table_cell_join"],
            source_fragments=[_fragment(text, 0)],
        )

        report = verify_element_text(self._document(element))

        assert report.verified == 0
        assert len(report.unprovable) == 1
        assert rebuild_row_text(element) is None

    def test_a_row_whose_cells_do_not_make_its_text_fails(self) -> None:
        """Two recorded facts disagreeing is a failure, not an exemption."""

        element = _row_document(1, 2, headers=None).elements[0]
        element.text = element.text + " and something the cells do not say"

        report = verify_element_text(self._document(element))

        assert report.unprovable == []
        assert len(report.failures) == 1
        assert element.element_id in report.failures[0]

    def test_a_row_opening_on_an_empty_cell_still_reproduces(self) -> None:
        """The one tolerance is the strip the element builder itself applies."""

        cells = ["", _value(1, 1)]
        text = ROW_SEPARATOR.join(cells).strip()
        element = CanonicalElement(
            element_id="E000000",
            element_type="table_row",
            logical_order=0,
            text=text,
            transformations=["table_cell_join"],
            source_fragments=[_fragment(text, 0)],
            table_structure=TableStructure(
                cells=[
                    TableCell(
                        text=value,
                        position=TableCellRef(row_index=1, column_index=column),
                    )
                    for column, value in enumerate(cells)
                ]
            ),
        )

        assert verify_element_text(self._document(element)).verified == 1


# ── the producers ────────────────────────────────────────────────────────


class TestAParserRecordsWhatItRead:
    """End to end from a real file, so the carry is not merely constructible."""

    @staticmethod
    def _ingest(tmp_path, grid: list[list[str]]):
        from docx import Document as DocxDocument

        document = DocxDocument()
        table = document.add_table(rows=len(grid), cols=len(grid[0]))
        for row_index, row in enumerate(grid):
            for column_index, value in enumerate(row):
                table.cell(row_index, column_index).text = value
        path = tmp_path / "grid.docx"
        document.save(str(path))
        return document_ingestion.ingest_docx(str(path), "grid")

    def test_a_parsed_row_carries_the_cells_it_was_joined_from(self, tmp_path) -> None:
        grid = [
            [_label(0), _label(1), _label(2)],
            [_value(1, 0), _value(1, 1), _value(1, 2)],
            [_value(2, 0), _value(2, 1), _value(2, 2)],
        ]

        document = self._ingest(tmp_path, grid)
        rows = [e for e in document.elements if e.element_type == "table_row"]

        assert len(rows) == 2
        for offset, element in enumerate(rows):
            assert element.table_structure is not None
            assert [
                cell.text for cell in element.table_structure.ordered_cells
            ] == grid[offset + 1]
            # The grid's own row index, not the position in what was left after
            # the header row was taken off the front.
            assert {
                cell.position.row_index for cell in element.table_structure.ordered_cells
            } == {offset + 1}

    def test_a_parsed_row_pairs_its_values_with_the_labels_the_grid_stated(
        self, tmp_path
    ) -> None:
        grid = [
            [_label(0), _label(1)],
            [_value(1, 0), _value(1, 1)],
        ]

        document = self._ingest(tmp_path, grid)
        row = next(e for e in document.elements if e.element_type == "table_row")
        assert row.table_structure is not None

        assert row.table_structure.header_value_pairs() == [
            ((_label(0),), row.table_structure.ordered_cells[0]),
            ((_label(1),), row.table_structure.ordered_cells[1]),
        ]

    def test_a_parsed_grid_that_stated_no_labels_records_none(self, tmp_path) -> None:
        """A headerless grid keeps its shape and gains no names.

        The grid repeats its first row's values further down, which is how the
        parser establishes that row 0 states values rather than labels: labels
        are distinct by construction and values recur. The verdict is the
        parser's, and the carrier records it rather than second-guessing it.
        """

        repeated = [_value(0, 0), _value(0, 1)]
        grid = [repeated, list(repeated), list(repeated)]

        document = self._ingest(tmp_path, grid)
        rows = [e for e in document.elements if e.element_type == "table_row"]

        assert len(rows) == len(grid), "the fixture's row 0 was read as a header"
        for offset, element in enumerate(rows):
            assert element.table_headers is None
            assert element.table_structure is not None
            assert element.table_structure.column_labels is None
            assert element.table_structure.labels_by_column() == {}
            assert [
                cell.text for cell in element.table_structure.ordered_cells
            ] == repeated
            assert {
                cell.position.row_index
                for cell in element.table_structure.ordered_cells
            } == {offset}

    def test_a_parsed_row_survives_the_clause_projection(self, tmp_path) -> None:
        """The whole chain, from a file on disk to a rebuilt document."""

        grid = [
            [_label(0), _label(1)],
            [_value(1, 0), _value(1, 1)],
        ]

        document = self._ingest(tmp_path, grid)
        rebuilt = _round_trip(document)
        row = next(e for e in rebuilt.elements if e.element_type == "table_row")

        assert row.table_structure is not None
        assert [cell.text for cell in row.table_structure.ordered_cells] == grid[1]
        assert [names for names, _ in row.table_structure.header_value_pairs()] == [
            (_label(0),),
            (_label(1),),
        ]


# ── the guard over the carrier's own code ────────────────────────────────

#: Subjects, documents, organisations and languages drawn from the corpora this
#: platform has been run against and from this suite's own fixtures. Listed here,
#: in a test, which is the only place any of them belongs.
DOMAIN_WORDS = (
    "ais",
    "hardware",
    "laptop",
    "leave",
    "leaves",
    "absence",
    "penalty",
    "penalties",
    "vacation",
    "handbook",
    "employee",
    "payroll",
    "salary",
    "allowance",
    "housing",
    "gmu",
    "arabic",
    "english",
    "maritime",
    "veterinary",
    "procurement",
)

#: The code that implements the carrier. Held to the rule directly rather than
#: through a description of it: a guard that reads a list of module names cannot
#: notice the day one of them stops being the implementation.
CARRIER_CODE = (
    canonical_rebuild,
    document_ingestion._row_structure,
    render_table_cells,
    rebuild_row_text,
    table_structure_of,
    table_cell_of,
    covered_columns,
    TableStructure,
    TableCell,
)


def _executable_source(obj) -> str:
    """The object's code with its docstrings and comments removed.

    Documentation is not code. A comment recording why a rule exists, including
    the case it was observed against, is provenance a maintainer needs; a literal
    a branch can turn on is a shape a corpus can steer. Only the second is
    checked here, which is what makes the distinction enforceable rather than
    merely stated.
    """

    tree = ast.parse(textwrap.dedent(inspect.getsource(obj)))
    for node in ast.walk(tree):
        if not isinstance(
            node, (ast.Module, ast.ClassDef, ast.FunctionDef, ast.AsyncFunctionDef)
        ):
            continue
        body = node.body
        if (
            body
            and isinstance(body[0], ast.Expr)
            and isinstance(body[0].value, ast.Constant)
            and isinstance(body[0].value.value, str)
        ):
            node.body = body[1:] or [ast.Pass()]
    return ast.unparse(tree)


@pytest.mark.parametrize("obj", CARRIER_CODE, ids=lambda o: getattr(o, "__name__", str(o)))
def test_the_carrier_names_no_subject_of_any_document(obj) -> None:
    """The record is generic or it is not a record, it is a fitting."""

    source = _executable_source(obj).lower()
    found = [
        word for word in DOMAIN_WORDS if re.search(rf"\b{re.escape(word)}\b", source)
    ]

    assert found == [], f"{getattr(obj, '__name__', obj)} names {found}"


@pytest.mark.parametrize("obj", CARRIER_CODE, ids=lambda o: getattr(o, "__name__", str(o)))
def test_the_carrier_carries_no_measured_magnitude(obj) -> None:
    """No count fitted to a grid somebody measured.

    A table's shape is read from the table. Nothing here may hold a number of
    rows, of columns or of anything else that came from looking at one, so the
    rule is drawn where it can be checked: the only integers this code is allowed
    to state are the small ones a boundary needs.
    """

    tree = ast.parse(_executable_source(obj))
    magnitudes = [
        node.value
        for node in ast.walk(tree)
        if isinstance(node, ast.Constant)
        and isinstance(node.value, (int, float))
        and not isinstance(node.value, bool)
        and abs(node.value) > 1
    ]

    assert magnitudes == [], f"{getattr(obj, '__name__', obj)} states {magnitudes}"
